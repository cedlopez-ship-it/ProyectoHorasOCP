import logging
import time
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from dataverse import get_horas_ingeniero, get_opp_por_guid, get_horas_equipo
from session_manager import cache_get, cache_opps, cambiar_estado
from webex import send_message, send_file

def texto_tipo(tipo):

    tipos = {
        "soporte_normal": "Soporte - Horas normales",
        "soporte_extra": "Soporte - Horas extras",
        "implementacion_normal": "Implementación - Horas normales",
        "implementacion_extra": "Implementación - Horas extras",
        "consultoria": "Consultoría",
        "consultoría": "Consultoría"
    }

    return tipos.get(str(tipo or "").lower(), tipo or "")

def fecha_excel(fecha):

    if not fecha:
        return ""

    if isinstance(fecha, datetime):
        return fecha.date()

    try:
        return datetime.fromisoformat(str(fecha).replace("Z", "+00:00")).date()
    except ValueError:
        return fecha

def generar_excel_registros(ingeniero, fecha_desde, fecha_hasta, registros):

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Registros"

    encabezados = [
        "Fecha", "Tipo", "Cliente", "OPP", "Proyecto",
        "Inicio", "Fin", "Horas", "Comentario"
    ]

    hoja.append(encabezados)

    for registro in registros:
        hoja.append([
            fecha_excel(registro.get("hxp_horasfecha")),
            texto_tipo(registro.get("hxp_horastipo")),
            registro.get("cliente", ""),
            registro.get("opp", ""),
            registro.get("proyecto", ""),
            registro.get("hxp_horainicio_", ""),
            registro.get("hxp_horafin_", ""),
            registro.get("hxp_horascantidaddecimal", 0),
            registro.get("hxp_comentario", "")
        ])

    color_encabezado = PatternFill("solid", fgColor="1F4E78")

    for celda in hoja[1]:
        celda.fill = color_encabezado
        celda.font = Font(bold=True, color="FFFFFF")

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(len(str(celda.value or "")) for celda in columna) + 2
        hoja.column_dimensions[letra].width = min(ancho, 50)

    for celda in hoja["A"][1:]:
        celda.number_format = "dd-mm-yyyy"

    resumen = libro.create_sheet("Resumen")

    horas_por_tipo = {
        "soporte_normal": 0,
        "soporte_extra": 0,
        "implementacion_normal": 0,
        "implementacion_extra": 0,
        "consultoria": 0
    }

    for registro in registros:
        tipo = str(registro.get("hxp_horastipo", "")).lower()
        if tipo == "consultoría":
            tipo = "consultoria"
        if tipo in horas_por_tipo:
            horas_por_tipo[tipo] += registro.get("hxp_horascantidaddecimal") or 0

    datos_resumen = [
        ["Ingeniero", ingeniero],
        [
            "Período consultado",
            f"{fecha_desde.strftime('%d-%m-%Y')} al {fecha_hasta.strftime('%d-%m-%Y')}"
        ],
        ["Cantidad de registros", len(registros)],
        ["Soporte - Horas normales", horas_por_tipo["soporte_normal"]],
        ["Soporte - Horas extras", horas_por_tipo["soporte_extra"]],
        ["Implementación - Horas normales", horas_por_tipo["implementacion_normal"]],
        ["Implementación - Horas extras", horas_por_tipo["implementacion_extra"]],
        ["Horas consultoría", horas_por_tipo["consultoria"]],
        ["Total de horas", sum(horas_por_tipo.values())],
        ["Fecha de generación", datetime.now().strftime("%d-%m-%Y %H:%M")]
    ]

    for fila in datos_resumen:
        resumen.append(fila)

    for celda in resumen["A"]:
        celda.fill = color_encabezado
        celda.font = Font(bold=True, color="FFFFFF")

    resumen.column_dimensions["A"].width = 25
    resumen.column_dimensions["B"].width = 42

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return archivo

def manejar_mis_registros(room, s, trace_id):

    try:
        ingeniero_data = s.get("ingeniero_data", {})
        ingeniero_id = ingeniero_data.get("hxp_ingenieroid")

        if not ingeniero_id:
            logging.error(f"[{trace_id}] Ingeniero sin GUID para Mis registros")
            send_message(room, "❌ No se pudo identificar al ingeniero.")
            return

        fecha_hasta = datetime.now().date()
        fecha_desde = fecha_hasta - timedelta(days=30)

        logging.info(f"[{trace_id}] Inicio consulta Mis registros")

        registros = get_horas_ingeniero(
            ingeniero_id,
            fecha_desde.isoformat(),
            fecha_hasta.isoformat(),
            trace_id
        )

        if registros is None:
            send_message(
                room,
                "❌ No se pudieron consultar tus registros. Intentá nuevamente más tarde."
            )
            return

        logging.info(f"[{trace_id}] Cantidad de registros: {len(registros)}")

        opps_consultadas = 0
        opps_resueltas = {}

        for registro in registros:
            opp_guid = registro.get("_hxp_horasopp_value")

            if opp_guid:
                opp_key = opp_guid.lower()

                if opp_key not in opps_resueltas:
                    cache_key = f"guid_{opp_key}"
                    if cache_get(cache_opps, cache_key) is None:
                        opps_consultadas += 1

                    opps_resueltas[opp_key] = get_opp_por_guid(opp_guid, trace_id) or {}

                opp = opps_resueltas[opp_key]
            else:
                opp = {}

            registro["cliente"] = opp.get("cliente", "")
            registro["opp"] = opp.get("codigo", "")
            registro["proyecto"] = opp.get("proyecto", "")

        logging.info(f"[{trace_id}] Cantidad de OPP consultadas: {opps_consultadas}")

        inicio_generacion = time.time()
        archivo = generar_excel_registros(
            ingeniero_data.get("hxp_ingenieronombre", ingeniero_data.get("hxp_ingenieromail", "")),
            fecha_desde,
            fecha_hasta,
            registros
        )

        logging.info(
            f"[{trace_id}] Tiempo de generación Excel: {round(time.time() - inicio_generacion, 2)}s"
        )

        inicio_envio = time.time()
        enviado = send_file(
            room,
            f"mis_registros_{fecha_hasta.strftime('%Y%m%d')}.xlsx",
            archivo,
            trace_id
        )

        logging.info(
            f"[{trace_id}] Tiempo de envío Excel: {round(time.time() - inicio_envio, 2)}s"
        )

        if enviado:
            cambiar_estado(s, "mis_registros_continuar", trace_id)
            send_message(
                room,
                "✅ Te envié el Excel con tus registros de los últimos 30 días.\n\n"
                "¿Desea cargar horas?\n\n1️⃣ Sí\n2️⃣ No\n3️⃣ 🚪 Salir"
            )
        else:
            send_message(room, "❌ No se pudo enviar el archivo. Intentá nuevamente más tarde.")

    except Exception:
        logging.exception(f"[{trace_id}] Error en Mis registros")
        send_message(room, "❌ Ocurrió un error al generar tus registros.")


def generar_excel_equipo(fecha_desde, fecha_hasta, registros):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Horas del equipo"

    encabezados = [
        "Ingeniero", "Fecha", "Tipo", "Cliente", "OPP", "Proyecto",
        "Inicio", "Fin", "Horas", "Comentario"
    ]
    hoja.append(encabezados)

    for registro in registros:
        ingeniero = registro.get(
            "_hxp_horasingeniero_value@OData.Community.Display.V1.FormattedValue",
            ""
        )

        hoja.append([
            ingeniero,
            fecha_excel(registro.get("hxp_horasfecha")),
            texto_tipo(registro.get("hxp_horastipo")),
            registro.get("cliente", ""),
            registro.get("opp", ""),
            registro.get("proyecto", ""),
            registro.get("hxp_horainicio_", ""),
            registro.get("hxp_horafin_", ""),
            registro.get("hxp_horascantidaddecimal", 0),
            registro.get("hxp_comentario", "")
        ])

    color_encabezado = PatternFill("solid", fgColor="1F4E78")

    for celda in hoja[1]:
        celda.fill = color_encabezado
        celda.font = Font(bold=True, color="FFFFFF")

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(len(str(celda.value or "")) for celda in columna) + 2
        hoja.column_dimensions[letra].width = min(ancho, 50)

    for celda in hoja["B"][1:]:
        celda.number_format = "dd-mm-yyyy"

    resumen = libro.create_sheet("Resumen")

    horas_por_tipo = {
        "soporte_normal": 0,
        "soporte_extra": 0,
        "implementacion_normal": 0,
        "implementacion_extra": 0,
        "consultoria": 0
    }

    for registro in registros:
        tipo = str(registro.get("hxp_horastipo", "")).lower()
        if tipo == "consultoría":
            tipo = "consultoria"

        if tipo in horas_por_tipo:
            horas_por_tipo[tipo] += (
                registro.get("hxp_horascantidaddecimal") or 0
            )

    if fecha_desde and fecha_hasta:
        periodo = (
            f"{fecha_desde.strftime('%d-%m-%Y')} al "
            f"{fecha_hasta.strftime('%d-%m-%Y')}"
        )
    else:
        periodo = "Todo el histórico"

    datos_resumen = [
        ["Consulta", "Horas del equipo"],
        ["Período consultado", periodo],
        ["Cantidad de registros", len(registros)],
        ["Soporte - Horas normales", horas_por_tipo["soporte_normal"]],
        ["Soporte - Horas extras", horas_por_tipo["soporte_extra"]],
        ["Implementación - Horas normales", horas_por_tipo["implementacion_normal"]],
        ["Implementación - Horas extras", horas_por_tipo["implementacion_extra"]],
        ["Horas consultoría", horas_por_tipo["consultoria"]],
        ["Total de horas", sum(horas_por_tipo.values())],
        ["Fecha de generación", datetime.now().strftime("%d-%m-%Y %H:%M")]
    ]

    for fila in datos_resumen:
        resumen.append(fila)

    for celda in resumen["A"]:
        celda.fill = color_encabezado
        celda.font = Font(bold=True, color="FFFFFF")

    resumen.column_dimensions["A"].width = 35
    resumen.column_dimensions["B"].width = 42

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return archivo


def generar_y_enviar_excel_equipo(
    room,
    fecha_desde,
    fecha_hasta,
    trace_id,
    nombre_archivo
):
    try:
        logging.info(
            f"[{trace_id}] Inicio consulta de horas del equipo"
        )

        registros = get_horas_equipo(
            fecha_desde.isoformat() if fecha_desde else None,
            fecha_hasta.isoformat() if fecha_hasta else None,
            trace_id
        )

        if registros is None:
            send_message(
                room,
                "❌ No se pudieron consultar las horas del equipo. "
                "Intentá nuevamente más tarde."
            )
            return False

        logging.info(
            f"[{trace_id}] Horas del equipo: {len(registros)} registros"
        )

        opps_resueltas = {}

        for registro in registros:
            opp_guid = registro.get("_hxp_horasopp_value")

            if opp_guid:
                opp_key = opp_guid.lower()

                if opp_key not in opps_resueltas:
                    opps_resueltas[opp_key] = (
                        get_opp_por_guid(opp_guid, trace_id) or {}
                    )

                opp = opps_resueltas[opp_key]
            else:
                opp = {}

            registro["cliente"] = opp.get("cliente", "")
            registro["opp"] = opp.get("codigo", "")
            registro["proyecto"] = opp.get("proyecto", "")

        archivo = generar_excel_equipo(
            fecha_desde,
            fecha_hasta,
            registros
        )

        enviado = send_file(
            room,
            nombre_archivo,
            archivo,
            trace_id
        )

        if not enviado:
            send_message(
                room,
                "❌ No se pudo enviar el archivo. Intentá nuevamente más tarde."
            )

        return enviado

    except Exception:
        logging.exception(
            f"[{trace_id}] Error generando Excel del equipo"
        )
        send_message(
            room,
            "❌ Ocurrió un error al generar el Excel del equipo."
        )
        return False
